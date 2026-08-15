from __future__ import annotations

import io
import os
from datetime import UTC, datetime
from uuid import UUID, uuid4

from docx import Document
from fastapi import FastAPI
from fastapi.testclient import TestClient
from pypdf import PdfReader

os.environ.setdefault("OPENAI_API_KEY", "test-openai-key")
os.environ.setdefault("DATABASE_URL", "postgresql://test:test@localhost:5432/test")

from app.github_summary.exports import (
    build_summary_docx,
    build_summary_pdf,
    export_filename,
)
from app.github_summary.models import (
    AnalysisStatistics,
    ApiEndpointSummary,
    ComponentSummary,
    RepositoryIdentity,
    RepositorySummary,
    RepositorySummaryResponse,
)
from app.routes.github import get_repository_summary_service, router


def _summary_response() -> RepositorySummaryResponse:
    return RepositorySummaryResponse(
        generated_at=datetime(2026, 8, 9, 21, 0, tzinfo=UTC),
        duration_seconds=12.5,
        repository=RepositoryIdentity(
            owner="example",
            name="support-api",
            full_name="example/support-api",
            url="https://github.com/example/support-api",
            visibility="public",
            description="Evidence-backed customer support service.",
            default_branch="main",
            analyzed_ref="feature/export",
            commit_sha="abc123",
            language_bytes={"Python": 1000, "SQL": 250},
            stars=3,
            forks=2,
            open_issues=1,
            archived=False,
        ),
        analysis=AnalysisStatistics(
            archive_bytes=1200,
            discovered_files=12,
            analyzed_files=10,
            skipped_files=2,
            analyzed_characters=5000,
            batches=2,
            cache_hit=False,
            summary_chunks=3,
            embedded_chunks=3,
        ),
        summary=RepositorySummary(
            title="Customer Support Intelligence",
            executive_summary="EXECUTIVE-COMPLETE summary text.",
            problem_statement="PROBLEM-COMPLETE problem text.",
            primary_capabilities=["CAPABILITY-COMPLETE"],
            architecture=["ARCHITECTURE-COMPLETE"],
            technology_stack=["TECHNOLOGY-COMPLETE"],
            key_components=[
                ComponentSummary(
                    name="COMPONENT-COMPLETE",
                    paths=["app/component.py"],
                    responsibility="RESPONSIBILITY-COMPLETE",
                )
            ],
            api_endpoints=[
                ApiEndpointSummary(
                    method="POST",
                    path="/complete-export",
                    purpose="ENDPOINT-PURPOSE-COMPLETE",
                    source_file="app/routes.py",
                )
            ],
            data_and_storage=["STORAGE-COMPLETE"],
            request_or_processing_flow=["FLOW-COMPLETE"],
            setup_and_run=["SETUP-COMPLETE"],
            strengths=["STRENGTH-COMPLETE"],
            risks_and_gaps=["RISK-COMPLETE"],
            recommended_next_steps=["NEXT-STEP-COMPLETE"],
            evidence_files=["EVIDENCE-COMPLETE.py"],
        ),
        summary_markdown="# Complete summary",
        model="test-model",
    )


def _docx_text(content: bytes) -> str:
    document = Document(io.BytesIO(content))
    values = [paragraph.text for paragraph in document.paragraphs]
    for table in document.tables:
        for row in table.rows:
            values.extend(cell.text for cell in row.cells)
    return "\n".join(values)


def test_docx_and_pdf_exports_include_complete_structured_summary() -> None:
    response = _summary_response()
    docx = build_summary_docx(response)
    pdf = build_summary_pdf(response)

    assert docx.startswith(b"PK")
    assert pdf.startswith(b"%PDF")
    docx_text = _docx_text(docx)
    pdf_text = "\n".join(page.extract_text() or "" for page in PdfReader(io.BytesIO(pdf)).pages)
    markers = {
        "EXECUTIVE-COMPLETE",
        "PROBLEM-COMPLETE",
        "CAPABILITY-COMPLETE",
        "ARCHITECTURE-COMPLETE",
        "TECHNOLOGY-COMPLETE",
        "COMPONENT-COMPLETE",
        "RESPONSIBILITY-COMPLETE",
        "/complete-export",
        "ENDPOINT-PURPOSE-COMPLETE",
        "STORAGE-COMPLETE",
        "FLOW-COMPLETE",
        "SETUP-COMPLETE",
        "STRENGTH-COMPLETE",
        "RISK-COMPLETE",
        "NEXT-STEP-COMPLETE",
        "EVIDENCE-COMPLETE.py",
    }
    assert all(marker in docx_text for marker in markers)
    assert all(marker in pdf_text for marker in markers)


def test_export_filename_is_safe_and_uses_real_word_extension() -> None:
    response = _summary_response()
    assert export_filename(response, "docx") == (
        "example-support-api-feature-export-repository-summary.docx"
    )
    assert export_filename(response, "pdf") == (
        "example-support-api-feature-export-repository-summary.pdf"
    )


def test_repository_summary_export_route_returns_downloadable_files() -> None:
    document_id = uuid4()

    class FakeService:
        async def response(self, requested_id: UUID):
            assert requested_id == document_id
            return _summary_response()

    app = FastAPI()
    app.include_router(router)
    app.dependency_overrides[get_repository_summary_service] = lambda: FakeService()
    client = TestClient(app)

    word = client.get(f"/api/v1/github/summary/{document_id}/export?format=docx")
    pdf = client.get(f"/api/v1/github/summary/{document_id}/export?format=pdf")

    assert word.status_code == 200
    assert word.content.startswith(b"PK")
    assert word.headers["content-type"].startswith(
        "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
    )
    assert ".docx" in word.headers["content-disposition"]
    assert pdf.status_code == 200
    assert pdf.content.startswith(b"%PDF")
    assert pdf.headers["content-type"].startswith("application/pdf")
    assert ".pdf" in pdf.headers["content-disposition"]