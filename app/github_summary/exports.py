from __future__ import annotations

import io
import re
from datetime import UTC
from xml.sax.saxutils import escape

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor
from reportlab.lib import colors
from reportlab.lib.enums import TA_LEFT
from reportlab.lib.pagesizes import letter
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.lib.units import inch
from reportlab.platypus import (
    KeepTogether,
    ListFlowable,
    ListItem,
    Paragraph,
    SimpleDocTemplate,
    Spacer,
    Table,
    TableStyle,
)

from app.github_summary.models import RepositorySummaryResponse


BLUE = "2E74B5"
DARK_BLUE = "1F4D78"
INK = "0B2545"
MUTED = "5F6B7A"
LIGHT_FILL = "E8EEF5"
BORDER = "D7DEE8"
NOT_IDENTIFIED = "Not identified from the analyzed repository files."


def export_filename(response: RepositorySummaryResponse, extension: str) -> str:
    stem = f"{response.repository.full_name}-{response.repository.analyzed_ref}-repository-summary"
    safe = re.sub(r"[^A-Za-z0-9._-]+", "-", stem).strip("-.")
    return f"{safe or 'repository-summary'}.{extension}"


def build_summary_docx(response: RepositorySummaryResponse) -> bytes:
    document = Document()
    section = document.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(1)
    section.right_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1)
    section.header_distance = Inches(0.492)
    section.footer_distance = Inches(0.492)
    _configure_docx_styles(document)
    _set_docx_header_footer(section, response)

    kicker = document.add_paragraph()
    kicker.paragraph_format.space_after = Pt(2)
    run = kicker.add_run("REPOSITORY SUMMARY")
    _set_run(run, size=10, bold=True, color=BLUE)

    title = document.add_paragraph(style="Title")
    title.add_run(response.summary.title)
    subtitle = document.add_paragraph(style="Subtitle")
    subtitle.add_run(
        f"{response.repository.full_name} | {response.repository.analyzed_ref}"
    )

    metadata = [
        ("Repository", response.repository.full_name),
        ("Repository URL", response.repository.url),
        ("Description", response.repository.description or NOT_IDENTIFIED),
        ("Default branch", response.repository.default_branch),
        ("Analyzed ref", response.repository.analyzed_ref),
        ("Commit SHA", response.repository.commit_sha or "Unknown"),
        ("Visibility", response.repository.visibility or NOT_IDENTIFIED),
        ("Generated", response.generated_at.astimezone(UTC).strftime("%Y-%m-%d %H:%M UTC")),
        ("Model", response.model),
        ("Files analyzed", f"{response.analysis.analyzed_files} of {response.analysis.discovered_files}"),
        ("Characters analyzed", f"{response.analysis.analyzed_characters:,}"),
        ("Summary batches", str(response.analysis.batches)),
    ]
    _add_docx_metadata_table(document, metadata)

    _add_docx_text_section(document, "Executive Summary", response.summary.executive_summary)
    _add_docx_text_section(document, "Problem Statement", response.summary.problem_statement)
    _add_docx_list_section(document, "Primary Capabilities", response.summary.primary_capabilities)
    _add_docx_list_section(document, "Architecture Overview", response.summary.architecture)
    _add_docx_list_section(document, "Technology Stack", response.summary.technology_stack)
    _add_docx_language_section(document, response)

    document.add_heading("Key Components", level=1)
    if response.summary.key_components:
        for component in response.summary.key_components:
            document.add_heading(component.name, level=2)
            _add_docx_body(document, component.responsibility)
            _add_docx_definition(document, "Paths", ", ".join(component.paths) or NOT_IDENTIFIED)
    else:
        _add_docx_body(document, NOT_IDENTIFIED)

    document.add_heading("API Endpoints", level=1)
    if response.summary.api_endpoints:
        for endpoint in response.summary.api_endpoints:
            document.add_heading(f"{endpoint.method.upper()} {endpoint.path}", level=2)
            _add_docx_body(document, endpoint.purpose)
            _add_docx_definition(document, "Source", endpoint.source_file or NOT_IDENTIFIED)
    else:
        _add_docx_body(document, NOT_IDENTIFIED)

    sections = (
        ("Data and Storage", response.summary.data_and_storage),
        ("Request and Processing Flow", response.summary.request_or_processing_flow),
        ("Setup and Run", response.summary.setup_and_run),
        ("Strengths", response.summary.strengths),
        ("Risks and Gaps", response.summary.risks_and_gaps),
        ("Recommended Next Steps", response.summary.recommended_next_steps),
        ("Evidence Files", response.summary.evidence_files),
    )
    for heading, items in sections:
        _add_docx_list_section(document, heading, items)

    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()


def build_summary_pdf(response: RepositorySummaryResponse) -> bytes:
    buffer = io.BytesIO()
    document = SimpleDocTemplate(
        buffer,
        pagesize=letter,
        leftMargin=inch,
        rightMargin=inch,
        topMargin=inch,
        bottomMargin=inch,
        title=response.summary.title,
        author="ReleaseMind",
        subject=f"Repository summary for {response.repository.full_name}",
    )
    styles = _pdf_styles()
    story = [
        Paragraph("REPOSITORY SUMMARY", styles["Kicker"]),
        Paragraph(escape(response.summary.title), styles["Title"]),
        Paragraph(
            escape(f"{response.repository.full_name} | {response.repository.analyzed_ref}"),
            styles["Subtitle"],
        ),
    ]
    metadata = [
        ("Repository", response.repository.full_name),
        ("Repository URL", response.repository.url),
        ("Description", response.repository.description or NOT_IDENTIFIED),
        ("Default branch", response.repository.default_branch),
        ("Analyzed ref", response.repository.analyzed_ref),
        ("Commit SHA", response.repository.commit_sha or "Unknown"),
        ("Visibility", response.repository.visibility or NOT_IDENTIFIED),
        ("Generated", response.generated_at.astimezone(UTC).strftime("%Y-%m-%d %H:%M UTC")),
        ("Model", response.model),
        ("Files analyzed", f"{response.analysis.analyzed_files} of {response.analysis.discovered_files}"),
        ("Characters analyzed", f"{response.analysis.analyzed_characters:,}"),
        ("Summary batches", str(response.analysis.batches)),
    ]
    story.extend([_pdf_metadata_table(metadata, styles), Spacer(1, 8)])
    _pdf_text_section(story, styles, "Executive Summary", response.summary.executive_summary)
    _pdf_text_section(story, styles, "Problem Statement", response.summary.problem_statement)
    _pdf_list_section(story, styles, "Primary Capabilities", response.summary.primary_capabilities)
    _pdf_list_section(story, styles, "Architecture Overview", response.summary.architecture)
    _pdf_list_section(story, styles, "Technology Stack", response.summary.technology_stack)
    languages = [
        f"{language}: {size:,} bytes"
        for language, size in sorted(
            response.repository.language_bytes.items(), key=lambda item: item[1], reverse=True
        )
    ]
    _pdf_list_section(story, styles, "Repository Languages", languages)

    story.append(Paragraph("Key Components", styles["Heading1"]))
    if response.summary.key_components:
        for component in response.summary.key_components:
            block = [
                Paragraph(escape(component.name), styles["Heading2"]),
                Paragraph(escape(component.responsibility), styles["Body"]),
                Paragraph(
                    f"<b>Paths:</b> {escape(', '.join(component.paths) or NOT_IDENTIFIED)}",
                    styles["Definition"],
                ),
            ]
            story.append(KeepTogether(block))
    else:
        story.append(Paragraph(NOT_IDENTIFIED, styles["Body"]))

    story.append(Paragraph("API Endpoints", styles["Heading1"]))
    if response.summary.api_endpoints:
        for endpoint in response.summary.api_endpoints:
            block = [
                Paragraph(
                    escape(f"{endpoint.method.upper()} {endpoint.path}"), styles["Heading2"]
                ),
                Paragraph(escape(endpoint.purpose), styles["Body"]),
                Paragraph(
                    f"<b>Source:</b> {escape(endpoint.source_file or NOT_IDENTIFIED)}",
                    styles["Definition"],
                ),
            ]
            story.append(KeepTogether(block))
    else:
        story.append(Paragraph(NOT_IDENTIFIED, styles["Body"]))

    sections = (
        ("Data and Storage", response.summary.data_and_storage),
        ("Request and Processing Flow", response.summary.request_or_processing_flow),
        ("Setup and Run", response.summary.setup_and_run),
        ("Strengths", response.summary.strengths),
        ("Risks and Gaps", response.summary.risks_and_gaps),
        ("Recommended Next Steps", response.summary.recommended_next_steps),
        ("Evidence Files", response.summary.evidence_files),
    )
    for heading, items in sections:
        _pdf_list_section(story, styles, heading, items)

    def page_furniture(canvas, _document) -> None:
        canvas.saveState()
        canvas.setFont("Helvetica", 8)
        canvas.setFillColor(colors.HexColor(f"#{MUTED}"))
        canvas.drawString(inch, letter[1] - 0.48 * inch, "ReleaseMind Repository Intelligence")
        canvas.drawRightString(
            letter[0] - inch,
            0.48 * inch,
            f"{response.repository.full_name} | Page {canvas.getPageNumber()}",
        )
        canvas.restoreState()

    document.build(story, onFirstPage=page_furniture, onLaterPages=page_furniture)
    return buffer.getvalue()


def _configure_docx_styles(document: Document) -> None:
    styles = document.styles
    _configure_style(styles["Normal"], 11, "000000", 0, 6, 1.25)
    _configure_style(styles["Title"], 24, INK, 0, 6, 1.0, bold=True)
    _configure_style(styles["Subtitle"], 13, MUTED, 0, 16, 1.0)
    _configure_style(styles["Heading 1"], 16, BLUE, 18, 10, 1.0, bold=True)
    _configure_style(styles["Heading 2"], 13, BLUE, 14, 7, 1.0, bold=True)
    _configure_style(styles["Heading 3"], 12, DARK_BLUE, 10, 5, 1.0, bold=True)
    _configure_style(styles["List Bullet"], 11, "000000", 0, 4, 1.25)
    styles["List Bullet"].paragraph_format.left_indent = Inches(0.375)
    styles["List Bullet"].paragraph_format.first_line_indent = Inches(-0.188)


def _configure_style(style, size, color, before, after, line_spacing, bold=False) -> None:
    style.font.name = "Calibri"
    style._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
    style._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
    style.font.size = Pt(size)
    style.font.bold = bold
    style.font.color.rgb = RGBColor.from_string(color)
    style.paragraph_format.space_before = Pt(before)
    style.paragraph_format.space_after = Pt(after)
    style.paragraph_format.line_spacing = line_spacing


def _set_run(run, *, size=None, bold=None, color=None) -> None:
    run.font.name = "Calibri"
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), "Calibri")
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), "Calibri")
    if size is not None:
        run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    if color is not None:
        run.font.color.rgb = RGBColor.from_string(color)


def _set_docx_header_footer(section, response: RepositorySummaryResponse) -> None:
    header = section.header.paragraphs[0]
    header.text = "ReleaseMind Repository Intelligence"
    header.style = "Normal"
    header.paragraph_format.space_after = Pt(0)
    for run in header.runs:
        _set_run(run, size=8, color=MUTED)

    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = footer.add_run(f"{response.repository.full_name} | Page ")
    _set_run(run, size=8, color=MUTED)
    field = OxmlElement("w:fldSimple")
    field.set(qn("w:instr"), "PAGE")
    footer._p.append(field)


def _add_docx_metadata_table(document: Document, rows: list[tuple[str, str]]) -> None:
    table = document.add_table(rows=0, cols=2)
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    table.autofit = False
    for label, value in rows:
        cells = table.add_row().cells
        cells[0].vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        cells[1].vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        label_paragraph = cells[0].paragraphs[0]
        label_paragraph.paragraph_format.space_after = Pt(0)
        label_run = label_paragraph.add_run(label)
        _set_run(label_run, size=9, bold=True, color=INK)
        value_paragraph = cells[1].paragraphs[0]
        value_paragraph.paragraph_format.space_after = Pt(0)
        value_run = value_paragraph.add_run(value)
        _set_run(value_run, size=9, color="000000")
        _shade_cell(cells[0], LIGHT_FILL)
        for cell in cells:
            _set_cell_margins(cell, top=80, bottom=80, start=120, end=120)
    _set_table_geometry(table, (2700, 6660))
    document.add_paragraph().paragraph_format.space_after = Pt(0)


def _set_table_geometry(table, widths: tuple[int, ...]) -> None:
    table_properties = table._tbl.tblPr
    table_width = table_properties.first_child_found_in("w:tblW")
    table_width.set(qn("w:w"), str(sum(widths)))
    table_width.set(qn("w:type"), "dxa")
    indent = OxmlElement("w:tblInd")
    indent.set(qn("w:w"), "120")
    indent.set(qn("w:type"), "dxa")
    table_properties.append(indent)
    grid = table._tbl.tblGrid
    for child in list(grid):
        grid.remove(child)
    for width in widths:
        column = OxmlElement("w:gridCol")
        column.set(qn("w:w"), str(width))
        grid.append(column)
    for row in table.rows:
        for cell, width in zip(row.cells, widths):
            cell.width = Inches(width / 1440)
            tc_width = cell._tc.get_or_add_tcPr().first_child_found_in("w:tcW")
            tc_width.set(qn("w:w"), str(width))
            tc_width.set(qn("w:type"), "dxa")


def _set_cell_margins(cell, **margins: int) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for margin, value in margins.items():
        node = tc_mar.find(qn(f"w:{margin}"))
        if node is None:
            node = OxmlElement(f"w:{margin}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def _shade_cell(cell, fill: str) -> None:
    shading = OxmlElement("w:shd")
    shading.set(qn("w:fill"), fill)
    cell._tc.get_or_add_tcPr().append(shading)


def _add_docx_body(document: Document, text: str) -> None:
    document.add_paragraph(text.strip() or NOT_IDENTIFIED, style="Normal")


def _add_docx_text_section(document: Document, heading: str, text: str) -> None:
    document.add_heading(heading, level=1)
    _add_docx_body(document, text)


def _add_docx_list_section(document: Document, heading: str, items: list[str]) -> None:
    document.add_heading(heading, level=1)
    cleaned = [item.strip() for item in items if item and item.strip()]
    if not cleaned:
        _add_docx_body(document, NOT_IDENTIFIED)
        return
    for item in cleaned:
        document.add_paragraph(item, style="List Bullet")


def _add_docx_definition(document: Document, label: str, value: str) -> None:
    paragraph = document.add_paragraph(style="Normal")
    label_run = paragraph.add_run(f"{label}: ")
    _set_run(label_run, bold=True, color=INK)
    paragraph.add_run(value)


def _add_docx_language_section(document: Document, response: RepositorySummaryResponse) -> None:
    items = [
        f"{language}: {size:,} bytes"
        for language, size in sorted(
            response.repository.language_bytes.items(), key=lambda item: item[1], reverse=True
        )
    ]
    _add_docx_list_section(document, "Repository Languages", items)


def _pdf_styles() -> dict[str, ParagraphStyle]:
    base = getSampleStyleSheet()
    return {
        "Kicker": ParagraphStyle(
            "Kicker", parent=base["Normal"], fontName="Helvetica-Bold", fontSize=10,
            textColor=colors.HexColor(f"#{BLUE}"), spaceAfter=2, leading=12,
        ),
        "Title": ParagraphStyle(
            "Title", parent=base["Title"], fontName="Helvetica-Bold", fontSize=24,
            textColor=colors.HexColor(f"#{INK}"), alignment=TA_LEFT, spaceAfter=6, leading=28,
        ),
        "Subtitle": ParagraphStyle(
            "Subtitle", parent=base["Normal"], fontName="Helvetica", fontSize=13,
            textColor=colors.HexColor(f"#{MUTED}"), spaceAfter=16, leading=16,
        ),
        "Body": ParagraphStyle(
            "Body", parent=base["BodyText"], fontName="Helvetica", fontSize=11,
            textColor=colors.black, spaceAfter=6, leading=13.75,
        ),
        "Definition": ParagraphStyle(
            "Definition", parent=base["BodyText"], fontName="Helvetica", fontSize=10,
            textColor=colors.HexColor(f"#{MUTED}"), spaceAfter=6, leading=12.5,
        ),
        "Heading1": ParagraphStyle(
            "Heading1", parent=base["Heading1"], fontName="Helvetica-Bold", fontSize=16,
            textColor=colors.HexColor(f"#{BLUE}"), spaceBefore=18, spaceAfter=10, leading=19,
            keepWithNext=True,
        ),
        "Heading2": ParagraphStyle(
            "Heading2", parent=base["Heading2"], fontName="Helvetica-Bold", fontSize=13,
            textColor=colors.HexColor(f"#{BLUE}"), spaceBefore=14, spaceAfter=7, leading=16,
            keepWithNext=True,
        ),
        "Bullet": ParagraphStyle(
            "Bullet", parent=base["BodyText"], fontName="Helvetica", fontSize=11,
            textColor=colors.black, spaceAfter=4, leading=13.75, leftIndent=0,
        ),
        "MetaLabel": ParagraphStyle(
            "MetaLabel", parent=base["Normal"], fontName="Helvetica-Bold", fontSize=9,
            textColor=colors.HexColor(f"#{INK}"), leading=11,
        ),
        "MetaValue": ParagraphStyle(
            "MetaValue", parent=base["Normal"], fontName="Helvetica", fontSize=9,
            textColor=colors.black, leading=11,
        ),
    }


def _pdf_metadata_table(rows: list[tuple[str, str]], styles) -> Table:
    data = [
        [Paragraph(escape(label), styles["MetaLabel"]), Paragraph(escape(value), styles["MetaValue"])]
        for label, value in rows
    ]
    table = Table(data, colWidths=[1.875 * inch, 4.625 * inch], hAlign="LEFT", repeatRows=0)
    table.setStyle(
        TableStyle(
            [
                ("BACKGROUND", (0, 0), (0, -1), colors.HexColor(f"#{LIGHT_FILL}")),
                ("GRID", (0, 0), (-1, -1), 0.35, colors.HexColor(f"#{BORDER}")),
                ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
                ("LEFTPADDING", (0, 0), (-1, -1), 6),
                ("RIGHTPADDING", (0, 0), (-1, -1), 6),
                ("TOPPADDING", (0, 0), (-1, -1), 4),
                ("BOTTOMPADDING", (0, 0), (-1, -1), 4),
            ]
        )
    )
    return table


def _pdf_text_section(story, styles, heading: str, text: str) -> None:
    story.append(Paragraph(escape(heading), styles["Heading1"]))
    story.append(Paragraph(escape(text.strip() or NOT_IDENTIFIED), styles["Body"]))


def _pdf_list_section(story, styles, heading: str, items: list[str]) -> None:
    story.append(Paragraph(escape(heading), styles["Heading1"]))
    cleaned = [item.strip() for item in items if item and item.strip()]
    if not cleaned:
        story.append(Paragraph(NOT_IDENTIFIED, styles["Body"]))
        return
    story.append(
        ListFlowable(
            [ListItem(Paragraph(escape(item), styles["Bullet"])) for item in cleaned],
            bulletType="bullet",
            start="circle",
            leftIndent=27,
            bulletFontName="Helvetica",
            bulletFontSize=8,
            bulletOffsetY=1,
        )
    )